"""Engine adapters shared by the docboss benchmarks.

Each adapter takes a path and returns the extracted text (or, for `open`,
anything that proves the document was parsed). Adapters for external
programs run them as subprocesses; `tool_path` finds them on PATH or in the
directories named by `DOCBOSS_BENCH_TOOLS` (a `:`-separated list), which is
how locally built antiword and catdoc binaries are passed in.
"""

from __future__ import annotations

import os
import shutil
import subprocess
from collections.abc import Callable


def tool_path(name: str) -> str | None:
    for directory in os.environ.get("DOCBOSS_BENCH_TOOLS", "").split(":"):
        if not directory:
            continue
        candidate = os.path.join(directory, name)
        if os.access(candidate, os.X_OK):
            return candidate
    return shutil.which(name)


class ToolCrashed(RuntimeError):
    """An external engine died on a signal rather than exiting with an error."""


def run_tool(cmd: list[str], env: dict[str, str] | None = None) -> str:
    done = subprocess.run(cmd, capture_output=True, env=env, timeout=120)
    if done.returncode < 0:
        raise ToolCrashed(f"{cmd[0]} killed by signal {-done.returncode}")
    if done.returncode != 0:
        raise RuntimeError(f"{cmd[0]} exited {done.returncode}")
    return done.stdout.decode("utf-8", "replace")


def docboss_open(path: str) -> object:
    import docboss

    return docboss.Document(path).format


def docboss_text(path: str) -> str:
    import docboss

    return docboss.Document(path).extract_text(headers_footers=True, notes=True)


def docboss_markdown(path: str) -> str:
    import docboss

    return docboss.Document(path).extract_markdown()


def docboss_html(path: str) -> str:
    import docboss

    return docboss.Document(path).extract_html()


def python_docx_open(path: str) -> object:
    import docx

    return len(docx.Document(path).paragraphs)


def python_docx_text(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts += [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(parts)


def docx2python_text(path: str) -> str:
    from docx2python import docx2python

    with docx2python(path) as result:
        return result.text


def mammoth_text(path: str) -> str:
    import mammoth

    with open(path, "rb") as handle:
        return mammoth.extract_raw_text(handle).value


def mammoth_markdown(path: str) -> str:
    import mammoth

    with open(path, "rb") as handle:
        return mammoth.convert_to_markdown(handle).value


def mammoth_html(path: str) -> str:
    import mammoth

    with open(path, "rb") as handle:
        return mammoth.convert_to_html(handle).value


def docx2txt_text(path: str) -> str:
    import docx2txt

    return docx2txt.process(path)


def pandoc(path: str, target: str) -> str:
    import pypandoc

    return pypandoc.convert_file(path, target, format="docx")


def pandoc_text(path: str) -> str:
    return pandoc(path, "plain")


def pandoc_markdown(path: str) -> str:
    return pandoc(path, "gfm")


def pandoc_html(path: str) -> str:
    return pandoc(path, "html")


def antiword_text(path: str) -> str:
    binary = tool_path("antiword")
    if not binary:
        raise RuntimeError("antiword not found")
    env = dict(os.environ)
    resources = os.path.join(os.path.dirname(binary), "Resources")
    if os.path.isdir(resources):
        env["ANTIWORDHOME"] = resources
    return run_tool([binary, "-m", "UTF-8.txt", "-w", "0", path], env)


def catdoc_text(path: str) -> str:
    binary = tool_path("catdoc")
    if not binary:
        raise RuntimeError("catdoc not found")
    return run_tool([binary, "-d", "utf-8", "-w", path])


Adapter = Callable[[str], object]

DOCX_ENGINES: dict[str, dict[str, Adapter]] = {
    "docboss": {
        "open": docboss_open,
        "text": docboss_text,
        "markdown": docboss_markdown,
        "html": docboss_html,
    },
    "python-docx": {"open": python_docx_open, "text": python_docx_text},
    "docx2python": {"text": docx2python_text},
    "mammoth": {"text": mammoth_text, "markdown": mammoth_markdown, "html": mammoth_html},
    "docx2txt": {"text": docx2txt_text},
    "pandoc": {"text": pandoc_text, "markdown": pandoc_markdown, "html": pandoc_html},
}

DOC_ENGINES: dict[str, dict[str, Adapter]] = {
    "docboss": {"open": docboss_open, "text": docboss_text},
    "antiword": {"text": antiword_text},
    "catdoc": {"text": catdoc_text},
}


def available(engines: dict[str, dict[str, Adapter]]) -> dict[str, dict[str, Adapter]]:
    """The engines whose module or program is installed."""
    probes = {
        "python-docx": "docx",
        "docx2python": "docx2python",
        "mammoth": "mammoth",
        "docx2txt": "docx2txt",
        "pandoc": "pypandoc",
    }
    out = {}
    for name, ops in engines.items():
        if name in probes:
            try:
                __import__(probes[name])
            except ImportError:
                continue
        if name in ("antiword", "catdoc") and not tool_path(name):
            continue
        out[name] = ops
    return out
